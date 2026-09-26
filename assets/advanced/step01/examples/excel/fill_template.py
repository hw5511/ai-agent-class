"""Copy the sales template and fill it with sample data from a JSON file.

Usage:
    python fill_template.py [--template FILE] [--data FILE] [--output FILE]

JSON layout (see sample_data.json):
    header  : dict  -> 작성일, 담당자, 부서, 기간
    sales   : list  -> 날짜, 거래처명, 제품명, 카테고리, 수량, 단가, 할인율, 비고
    monthly : list  -> 월, 목표_매출, 실제_매출 (optional), 비고 (optional)

Only the template's designated input cells are written. Formula cells
(H, J, totals, 달성률) are left untouched. If a month has no 실제_매출,
it is computed from the sales rows (sum of 할인 후 매출 for that month).
"""

import argparse
import json
import shutil
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook

TEMPLATE = "매출데이터_템플릿.xlsx"
DATA = "sample_data.json"
OUTPUT = "매출데이터_샘플.xlsx"

SALES_SHEET = "매출 데이터"
MONTHLY_SHEET = "월별 요약"
FIRST_DATA_ROW = 5
MAX_DATA_ROWS = 30
CATEGORIES = {"식품", "의류", "전자기기", "화장품", "생활용품", "서비스", "기타"}

HEADER_CELLS = {"작성일": "B3", "담당자": "E3", "부서": "H3", "기간": "K3"}


def parse_date(text):
    return datetime.strptime(text, "%Y-%m-%d")


def fill_header(ws, header):
    for key, cell in HEADER_CELLS.items():
        value = header.get(key)
        if value is None:
            continue
        if key == "작성일":
            ws[cell] = parse_date(value)
            ws[cell].number_format = "yyyy-mm-dd"
        else:
            ws[cell] = value


def fill_sales(ws, sales):
    if len(sales) > MAX_DATA_ROWS:
        raise ValueError(
            f"Template holds {MAX_DATA_ROWS} rows but {len(sales)} sales rows were given"
        )
    for offset, row in enumerate(sales):
        r = FIRST_DATA_ROW + offset
        category = row.get("카테고리", "")
        if category and category not in CATEGORIES:
            raise ValueError(f"Row {offset + 1}: unknown category '{category}'")
        ws[f"B{r}"] = parse_date(row["날짜"])
        ws[f"B{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"] = row.get("거래처명", "")
        ws[f"D{r}"] = row.get("제품명", "")
        ws[f"E{r}"] = category
        ws[f"F{r}"] = row["수량"]
        ws[f"G{r}"] = row["단가"]
        ws[f"I{r}"] = row.get("할인율", 0)
        ws[f"K{r}"] = row.get("비고", "")


def actuals_by_month(sales):
    totals = defaultdict(float)
    for row in sales:
        month = parse_date(row["날짜"]).month
        net = row["수량"] * row["단가"] * (1 - row.get("할인율", 0))
        totals[f"{month}월"] += net
    return totals


def fill_monthly(ws, monthly, sales):
    computed = actuals_by_month(sales)
    labels = {ws[f"A{r}"].value: r for r in range(3, 15)}
    for item in monthly:
        r = labels.get(item["월"])
        if r is None:
            raise ValueError(f"Unknown month label '{item['월']}'")
        if "목표_매출" in item:
            ws[f"B{r}"] = item["목표_매출"]
        actual = item.get("실제_매출", computed.get(item["월"]))
        if actual is not None:
            ws[f"C{r}"] = round(actual)
        if item.get("비고"):
            ws[f"E{r}"] = item["비고"]


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--template", default=TEMPLATE)
    parser.add_argument("--data", default=DATA)
    parser.add_argument("--output", default=OUTPUT)
    args = parser.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    shutil.copyfile(args.template, args.output)
    wb = load_workbook(args.output)

    fill_header(wb[SALES_SHEET], data.get("header", {}))
    fill_sales(wb[SALES_SHEET], data.get("sales", []))
    fill_monthly(wb[MONTHLY_SHEET], data.get("monthly", []), data.get("sales", []))

    wb.save(args.output)
    print(f"Wrote {args.output}: {len(data.get('sales', []))} sales rows, "
          f"{len(data.get('monthly', []))} monthly rows")


if __name__ == "__main__":
    main()
