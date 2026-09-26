"""Copy the sales template and fill it with data from a JSON file.

Usage:
    python fill_template.py --data DATA.json --output OUTPUT.xlsx [--template FILE]

The template defaults to ../assets/매출데이터_템플릿.xlsx (relative to this
script). See ../references/example_data.json for the JSON layout:

    header  : dict  -> 작성일, 담당자, 부서, 기간
    sales   : list  -> 날짜, 거래처명, 제품명, 카테고리, 수량, 단가, 할인율, 비고
    monthly : list  -> 월, 목표_매출, 실제_매출 (optional), 비고 (optional)

Only the template's designated input cells (blue fill) are written. Formula
cells (H, J, totals, 달성률) are left untouched so the workbook keeps
recalculating on its own. If a month has no 실제_매출, it is computed from
the sales rows (sum of 할인 후 매출 for that month).

Exit code is non-zero and a message is printed to stderr when the data does
not fit the template (too many rows, unknown category, bad date format).
"""

import argparse
import json
import shutil
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = SCRIPT_DIR.parent / "assets" / "매출데이터_템플릿.xlsx"

SALES_SHEET = "매출 데이터"
MONTHLY_SHEET = "월별 요약"
FIRST_DATA_ROW = 5
MAX_DATA_ROWS = 30
CATEGORIES = ("식품", "의류", "전자기기", "화장품", "생활용품", "서비스", "기타")

HEADER_CELLS = {"작성일": "B3", "담당자": "E3", "부서": "H3", "기간": "K3"}


def parse_date(text, where):
    try:
        return datetime.strptime(str(text), "%Y-%m-%d")
    except ValueError:
        raise ValueError(f"{where}: date '{text}' must be YYYY-MM-DD") from None


def require_number(value, where):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{where}: expected a number, got {value!r}")
    return value


def fill_header(ws, header):
    for key, cell in HEADER_CELLS.items():
        value = header.get(key)
        if value in (None, ""):
            continue
        if key == "작성일":
            ws[cell] = parse_date(value, "header.작성일")
            ws[cell].number_format = "yyyy-mm-dd"
        else:
            ws[cell] = value


def fill_sales(ws, sales):
    if len(sales) > MAX_DATA_ROWS:
        raise ValueError(
            f"template holds {MAX_DATA_ROWS} sales rows but {len(sales)} were given; "
            "split the data by period or trim it"
        )
    for offset, row in enumerate(sales):
        r = FIRST_DATA_ROW + offset
        where = f"sales[{offset}]"
        category = row.get("카테고리", "")
        if category and category not in CATEGORIES:
            raise ValueError(
                f"{where}: category '{category}' is not in the dropdown list "
                f"{', '.join(CATEGORIES)}"
            )
        ws[f"B{r}"] = parse_date(row["날짜"], where)
        ws[f"B{r}"].number_format = "yyyy-mm-dd"
        ws[f"C{r}"] = row.get("거래처명", "")
        ws[f"D{r}"] = row.get("제품명", "")
        ws[f"E{r}"] = category
        ws[f"F{r}"] = require_number(row["수량"], where + ".수량")
        ws[f"G{r}"] = require_number(row["단가"], where + ".단가")
        discount = require_number(row.get("할인율", 0), where + ".할인율")
        if not 0 <= discount <= 1:
            raise ValueError(
                f"{where}: 할인율 must be a fraction between 0 and 1 (10% -> 0.1)"
            )
        ws[f"I{r}"] = discount
        ws[f"K{r}"] = row.get("비고", "")


def actuals_by_month(sales):
    totals = defaultdict(float)
    for offset, row in enumerate(sales):
        month = parse_date(row["날짜"], f"sales[{offset}]").month
        net = row["수량"] * row["단가"] * (1 - row.get("할인율", 0))
        totals[f"{month}월"] += net
    return totals


def fill_monthly(ws, monthly, sales):
    computed = actuals_by_month(sales)
    labels = {ws[f"A{r}"].value: r for r in range(3, 15)}
    for offset, item in enumerate(monthly):
        where = f"monthly[{offset}]"
        r = labels.get(item.get("월"))
        if r is None:
            raise ValueError(
                f"{where}: month label must be one of {', '.join(labels)}"
            )
        if item.get("목표_매출") is not None:
            ws[f"B{r}"] = require_number(item["목표_매출"], where + ".목표_매출")
        actual = item.get("실제_매출", computed.get(item["월"]))
        if actual is not None:
            ws[f"C{r}"] = round(require_number(actual, where + ".실제_매출"))
        if item.get("비고"):
            ws[f"E{r}"] = item["비고"]


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--data", required=True, help="JSON data file")
    parser.add_argument("--output", required=True, help="xlsx file to create")
    parser.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    args = parser.parse_args()

    with open(args.data, encoding="utf-8") as f:
        data = json.load(f)

    shutil.copyfile(args.template, args.output)
    wb = load_workbook(args.output)
    try:
        fill_header(wb[SALES_SHEET], data.get("header", {}))
        fill_sales(wb[SALES_SHEET], data.get("sales", []))
        fill_monthly(wb[MONTHLY_SHEET], data.get("monthly", []), data.get("sales", []))
    except (ValueError, KeyError) as exc:
        Path(args.output).unlink(missing_ok=True)
        print(f"error: {exc}", file=sys.stderr)
        sys.exit(1)

    wb.save(args.output)
    print(
        f"wrote {args.output}: {len(data.get('sales', []))} sales rows, "
        f"{len(data.get('monthly', []))} monthly rows"
    )


if __name__ == "__main__":
    main()
