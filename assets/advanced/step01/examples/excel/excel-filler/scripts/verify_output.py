"""Print a plain-text summary of a filled workbook so the result can be checked
without opening Excel. Formula cells are recomputed here from the inputs, since
openpyxl cannot evaluate them.

Usage:
    python verify_output.py OUTPUT.xlsx
"""

import sys

from openpyxl import load_workbook


def main(path):
    wb = load_workbook(path)
    ws = wb["매출 데이터"]
    print(
        "header:", ws["B3"].value, "|", ws["E3"].value, "|",
        ws["H3"].value, "|", ws["K3"].value,
    )
    gross = net = 0.0
    filled = 0
    for r in range(5, 35):
        qty, price, disc = ws[f"F{r}"].value, ws[f"G{r}"].value, ws[f"I{r}"].value
        if isinstance(qty, (int, float)) and isinstance(price, (int, float)):
            filled += 1
            gross += qty * price
            net += qty * price * (1 - (disc or 0))
    print(f"sales rows filled: {filled}/30")
    print(f"expected H35 (매출액 합계):      {gross:,.0f}")
    print(f"expected J35 (할인 후 매출 합계): {net:,.0f}")
    if gross:
        print(f"expected I35 (평균 할인율):      {1 - net / gross:.1%}")
    ws2 = wb["월별 요약"]
    print("monthly (월, 목표, 실제, 달성률):")
    for r in range(3, 15):
        target, actual = ws2[f"B{r}"].value, ws2[f"C{r}"].value
        rate = f"{actual / target:.1%}" if target and actual else "-"
        print(f"  {ws2[f'A{r}'].value:>4} {str(target or '-'):>14} {str(actual or '-'):>14} {rate:>8}")


if __name__ == "__main__":
    main(sys.argv[1])
