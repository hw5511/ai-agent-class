# -*- coding: utf-8 -*-
"""Generator for the automation course session 1 student kit (practice A: vendor checker).

Deterministic: fixed seed, re-running yields byte-identical CSV/JSON.
Numbers listed in s1_excluded_hashes.json (found registered by verify_unregistered.py)
are skipped and redrawn.
"""
import hashlib
import csv, json, random, zipfile, io, datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
KIT = ROOT / "assets/automation/step01/kit"
ZIP = ROOT / "assets/automation/step01/거래처검증_실습키트.zip"
EXCL = Path(__file__).resolve().parent / "s1_excluded_hashes.json"
SEED = 20261022
N = 200

PREFIX = ["한빛", "다온", "새솔", "누리", "가온", "우람", "하늘", "푸른", "미래", "도담", "아름", "청솔", "해맞이", "늘품",
          "대성", "동원", "신화", "태광", "삼진", "금강", "세움", "라온", "보람", "한결", "예솔", "나래", "구름", "바른",
          "유진", "서린", "참빛", "다솜", "이든", "가람", "솔바람", "온누리", "별빛", "단비", "오름", "한울"]
SUFFIX = ["정밀", "식품", "물류", "건설", "산업", "상사", "테크", "시스템", "유통", "기계", "전자", "소프트", "푸드",
          "로지스", "엔지니어링", "패키지", "종합상사", "개발", "솔루션", "디자인", "화학", "금속", "운수", "서비스"]
SURN = ["김", "이", "박", "최", "정", "강", "조", "윤", "장", "임", "한", "오", "서", "신", "권", "황", "안", "송", "류", "홍"]
GIVEN = ["민준", "서연", "지훈", "하은", "도현", "수빈", "예준", "지우", "현우", "소영", "성민", "은지", "태양", "나연",
         "재원", "미란", "동혁", "유나", "상철", "보람", "경수", "혜진", "준호", "다혜", "영훈", "세진", "병철", "아영", "승우", "지원"]
BIZ = ["제조", "도소매", "운수", "서비스", "건설", "정보통신", "음식점"]
DEPT = ["구매팀", "총무팀", "영업지원팀", "생산관리팀"]


def check_digit(d9):
    w = [1, 3, 7, 1, 3, 7, 1, 3, 5]
    s = sum(d9[i] * w[i] for i in range(9)) + (d9[8] * 5) // 10
    return (10 - s % 10) % 10


def make_number(i, corp, excluded, used):
    for attempt in range(1000):
        r = random.Random(f"{SEED}-num-{i}-{attempt}")
        office = r.randint(101, 899)
        mid = r.choice([81, 86, 87]) if corp else r.randint(1, 79)
        serial = r.randint(0, 9999)
        d9 = [int(c) for c in f"{office:03d}{mid:02d}{serial:04d}"]
        num = "".join(map(str, d9)) + str(check_digit(d9))
        if hashlib.sha256(num.encode()).hexdigest() not in excluded and num not in used:
            return num
    raise RuntimeError("no number")


def fmt(num, kind):
    h = f"{num[:3]}-{num[3:5]}-{num[5:]}"
    return {"h": h, "p": num, "ht": h + " ", "pl": " " + num}[kind]


def build_rows():
    excluded = set(json.loads(EXCL.read_text(encoding="utf-8"))) if EXCL.exists() else set()
    g = random.Random(SEED)
    # one pair per prefix would give only 40 names; shuffle all 960 combos, keep distinct full names
    pairs = [(p, s) for p in PREFIX for s in SUFFIX]
    g.shuffle(pairs)
    names = pairs[:N]
    kinds = ["h"] * 140 + ["p"] * 50 + ["ht"] * 5 + ["pl"] * 5
    g.shuffle(kinds)
    rep_seen = set()
    rows, used = [], set()
    for i in range(N):
        p, s = names[i]
        corp = g.random() < 0.45
        name = f"(주){p}{s}" if corp else f"{p}{s}"
        while True:
            rep = g.choice(SURN) + g.choice(GIVEN)
            if rep not in rep_seen:
                rep_seen.add(rep)
                break
        num = make_number(i, corp, excluded, used)
        used.add(num)
        rows.append([f"V{i+1:04d}", name, rep, fmt(num, kinds[i]), g.choice(BIZ), g.choice(DEPT)])
    return rows


def rand_date(r, lo, hi):
    a = datetime.date(*lo)
    b = datetime.date(*hi)
    d = a + datetime.timedelta(days=r.randint(0, (b - a).days))
    return d.strftime("%Y%m%d")


def entry(b_no, stt="", stt_cd="", tax="", tax_cd="", end="", utcc="", rbf="", rbf_cd=""):
    return {"b_no": b_no, "b_stt": stt, "b_stt_cd": stt_cd, "tax_type": tax, "tax_type_cd": tax_cd,
            "end_dt": end, "utcc_yn": utcc, "tax_type_change_dt": "", "invoice_apply_dt": "",
            "rbf_tax_type": rbf, "rbf_tax_type_cd": rbf_cd}


def build_samples(rows):
    r = random.Random(SEED + 1)
    labels = ["gen"] * 150 + ["simple"] * 16 + ["free"] * 8 + ["rest"] * 8 + ["closed"] * 12 + ["none"] * 6
    r.shuffle(labels)
    cont = [i for i, l in enumerate(labels) if l in ("gen", "simple", "free")]
    chg_idx = set(r.sample(cont, 5))
    data = []
    for i, lab in enumerate(labels):
        b = "".join(c for c in rows[i][3] if c.isdigit())
        tail = dict(utcc="N", rbf="해당없음", rbf_cd="99")
        if lab == "gen":
            e = entry(b, "계속사업자", "01", "부가가치세 일반과세자", "01", **tail)
        elif lab == "simple":
            e = entry(b, "계속사업자", "01", "부가가치세 간이과세자", "02", **tail)
        elif lab == "free":
            e = entry(b, "계속사업자", "01", "부가가치세 면세사업자", "04", **tail)
        elif lab == "rest":
            e = entry(b, "휴업자", "02", "부가가치세 일반과세자", "01", **tail)
        elif lab == "closed":
            e = entry(b, "폐업자", "03", "폐업자", "", end=rand_date(r, (2024, 1, 1), (2026, 9, 30)), **tail)
        else:
            e = entry(b, tax="국세청에 등록되지 않은 사업자등록번호입니다.")
        if i in chg_idx:
            e["tax_type_change_dt"] = rand_date(r, (2026, 1, 1), (2026, 9, 30))
        data.append(e)
    return data


def write(path, text, bom=False):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig" if bom else "utf-8", newline="") as f:
        f.write(text)


def make_xlsx(path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.formatting.rule import CellIsRule
    wb = Workbook()
    ws = wb.active
    ws.title = "검증결과"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "거래처 사업자 상태 검증 결과"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "조회일"
    ws["A2"].font = Font(bold=True)
    ws["B2"].border = border
    for col, lab in zip("ACEG", ["전체", "정상", "주의", "위험"]):
        ws[f"{col}3"] = lab
        ws[f"{col}3"].font = Font(bold=True)
        ws[f"{col}3"].border = border
        ws[f"{chr(ord(col) + 1)}3"].border = border
    heads = ["거래처코드", "상호", "대표자", "사업자등록번호", "사업자상태", "과세유형", "폐업일", "판정", "사유"]
    widths = [12, 22, 12, 18, 14, 22, 12, 10, 40]
    fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    for c, (h, wd) in enumerate(zip(heads, widths), start=1):
        cell = ws.cell(row=6, column=c, value=h)
        cell.fill = fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = wd
    for r in range(7, 501):
        ws.cell(row=r, column=4).number_format = "@"
    ws.freeze_panes = "A7"
    ws.conditional_formatting.add("H7:H500", CellIsRule(operator="equal", formula=['"위험"'],
        fill=PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")))
    ws.conditional_formatting.add("H7:H500", CellIsRule(operator="equal", formula=['"주의"'],
        fill=PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")))
    wb.save(path)


API_MD = """# 국세청 사업자등록 상태조회 API 안내

주소는 https://api.odcloud.kr/api/nts-businessman/v1/status 이고 POST 방식으로 호출합니다.

인증키는 주소 뒤에 serviceKey 쿼리 파라미터로 붙입니다. 공공데이터포털의 일반 인증키 중 Decoding 값을 쓰려면 URL 인코딩해서 넣고, Encoding 값은 이미 인코딩되어 있으니 그대로 넣습니다. 키는 .env 파일에만 두고 코드나 채팅에 붙여 넣지 않습니다.

요청 헤더는 Content-Type: application/json 입니다. 본문은 JSON이고 b_no 배열에 사업자등록번호를 담습니다.

{"b_no": ["1234567890", "2345678901"]}

한 번에 최대 100개까지 조회할 수 있어서 200개면 두 번 나눠 호출합니다. 100개를 넘겨 보내면 413 TOO_LARGE_REQUEST 오류가 납니다. 하루 호출 한도는 개발계정 기준 100만 건입니다. 번호는 하이픈과 공백 없이 숫자 10자리만 보냅니다. 엑셀에 섞여 있는 하이픈이나 빈 칸은 보내기 전에 지워야 합니다.

응답은 request_cnt(요청 수), match_cnt(국세청에 등록된 건수), status_code(OK면 정상), data(번호별 결과 배열)로 구성됩니다. data의 각 항목은 다음 필드를 가집니다.

b_no: 조회한 사업자등록번호 (숫자만)
b_stt: 납세자 상태 이름 (계속사업자, 휴업자, 폐업자)
b_stt_cd: 납세자 상태 코드
tax_type: 과세유형 이름
tax_type_cd: 과세유형 코드
end_dt: 폐업일 (YYYYMMDD)
utcc_yn: 단위과세 전환 폐업 여부 (Y 또는 N)
tax_type_change_dt: 최근 과세유형 전환일 (YYYYMMDD)
invoice_apply_dt: 세금계산서 적용일 (YYYYMMDD)
rbf_tax_type: 직전 과세유형 이름
rbf_tax_type_cd: 직전 과세유형 코드

b_stt_cd 값은 01이 계속사업자, 02가 휴업자, 03이 폐업자입니다.

tax_type_cd 값은 다음과 같습니다(공공데이터포털 명세 기준).

01: 부가가치세 일반과세자
02: 부가가치세 간이과세자
03: 부가가치세 과세특례자
04: 부가가치세 면세사업자
05: 수익사업을 영위하지 않는 비영리법인이거나 고유번호가 부여된 단체, 국가기관 등
06: 고유번호가 부여된 단체
07: 부가가치세 간이과세자(세금계산서 발급사업자)

직전 과세유형이 없으면 rbf_tax_type_cd는 99(해당없음)로 옵니다. 국세청 자료는 30분마다 갱신되고, 새로 개업한 사업자는 조회되기까지 1~2일이 걸립니다.

국세청에 등록되지 않은 번호는 오류가 아니라 정상 응답으로 오며, b_stt_cd가 빈 문자열이고 tax_type에 안내 문구가 들어 있습니다.
"""

START_MD = """# 시작하기

이 폴더는 거래처 검증기 실습(실습 A)에 쓰는 재료입니다. 파일마다 역할이 다릅니다.

vendors.csv는 가상의 거래처 200곳 목록입니다. 사업자등록번호 표기가 하이픈, 숫자만, 공백 섞임으로 제각각인데, 실제 업무 데이터처럼 일부러 그렇게 만들었습니다.
vendors.csv의 상호, 대표자, 번호는 모두 지어낸 것입니다. 그래서 실제 국세청 API로 조회하면 전부 "등록되지 않은 사업자등록번호"로 응답합니다.
sample_raw 폴더에는 같은 형식의 연습용 응답 파일이 두 개 들어 있습니다(batch_001.json은 앞 100곳, batch_002.json은 뒤 100곳). 정상, 휴업, 폐업, 미등록이 섞여 있어서 판정 실습에 씁니다.
거래처검증_템플릿.xlsx는 서식이 모두 만들어진 결과 시트입니다. 스크립트는 값만 채우고 서식은 건드리지 않습니다.
API_안내.md는 국세청 API 호출 방법과 응답 필드, 코드표를 정리한 문서입니다.
.env.example은 인증키를 넣는 자리입니다. 파일 이름을 .env로 바꾸고 본인의 일반 인증키(Decoding)를 붙여 넣으세요.

작업은 Claude Code에게 이 폴더를 열어 보여 주고, 하고 싶은 일을 말로 설명하며 진행합니다. 코드는 직접 쓰지 않습니다.
"""


def main():
    rows = build_rows()
    KIT.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["거래처코드", "상호", "대표자", "사업자등록번호", "업종", "담당부서"])
    w.writerows(rows)
    write(KIT / "vendors.csv", buf.getvalue(), bom=True)

    data = build_samples(rows)
    for n in (1, 2):
        part = data[(n - 1) * 100: n * 100]
        obj = {"request_cnt": 100, "match_cnt": sum(1 for e in part if e["b_stt_cd"]),
               "status_code": "OK", "data": part}
        write(KIT / f"sample_raw/batch_{n:03d}.json", json.dumps(obj, ensure_ascii=False, indent=2) + "\n")

    make_xlsx(KIT / "거래처검증_템플릿.xlsx")
    write(KIT / "API_안내.md", API_MD)
    write(KIT / ".env.example", "NTS_SERVICE_KEY=여기에_일반_인증키(Decoding)를_붙여넣으세요")
    write(KIT / "시작하기.md", START_MD)

    files = sorted(p for p in KIT.rglob("*") if p.is_file())
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in files:
            zi = zipfile.ZipInfo("거래처검증_실습키트/" + p.relative_to(KIT).as_posix(), (2026, 10, 22, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, p.read_bytes())
    print("rows", len(rows), "files", len(files))


if __name__ == "__main__":
    main()
