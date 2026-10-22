# -*- coding: utf-8 -*-
"""Generator for the automation course session 1 alternative practice kit (travel expense settlement).

Deterministic: all expense rows are hand-written below (no randomness); re-running gives a
byte-identical expenses.csv and zip. sample_raw/ is real ECOS data recorded once with the keyless
`sample` key; it is fetched only if missing (or with --refetch) and then reused untouched.

Before writing expenses.csv the planted over-limit groups are verified against the real rates.
"""
import csv
import io
import json
import re
import sys
import time
import urllib.request
import zipfile
import datetime as dt
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
KIT = ROOT / "assets/automation/step01/expense_kit"
ZIP = ROOT / "assets/automation/step01/출장경비_실습키트.zip"
TOP = "출장경비_실습키트"
RAW = KIT / "sample_raw"
FETCH_DATE = "2026-10-22"          # day the sample_raw responses were recorded
RANGE = ("20260825", "20261016")
CODES = {"USD": "0000001", "JPY": "0000002", "EUR": "0000003", "CNY": "0000053", "VND": "0000035"}
UNIT = {"USD": 1, "EUR": 1, "CNY": 1, "JPY": 100, "VND": 100}
LIMIT = {"식비": 100000, "숙박": 250000, "기타": 50000}   # 교통: no limit
ZIP_TIME = (2026, 10, 22, 0, 0, 0)
HOLIDAYS = {"2026-09-24", "2026-09-25", "2026-09-26", "2026-10-03", "2026-10-05", "2026-10-09"}

# ---------------------------------------------------------------- expense data
# (date, item, amount, evidence, memo[, "c" -> write amount with thousands separator])
TRIPS = [
    dict(name="박민석", dept="해외영업팀", city="뉴욕", country="미국", cur="USD", rows=[
        ("2026-09-07", "교통", "68.50", "Y", "공항 택시"),
        ("2026-09-07", "숙박", "175.00", "Y", ""),
        ("2026-09-07", "식비", "24.80", "Y", "도착 후 저녁"),
        ("2026-09-08", "식비", "14.50", "Y", "조식"),
        ("2026-09-08", "식비", "22.30", "Y", ""),
        ("2026-09-08", "식비", "96.40", "Y", "거래처 저녁"),
        ("2026-09-08", "숙박", "175.00", "Y", ""),
        ("2026-09-08", "교통", "34.00", "Y", "지하철 정액권"),
        ("2026-09-09", "식비", "18.90", "Y", ""),
        ("2026-09-09", "식비", "41.20", "Y", "저녁"),
        ("2026-09-09", "숙박", "214.00", "Y", "성수기 요금"),
        ("2026-09-09", "교통", "27.60", "Y", "우버"),
        ("2026-09-09", "기타", "35.00", "Y", "전시회 입장권"),
        ("2026-09-10", "식비", "15.20", "N", "조식"),
        ("2026-09-10", "식비", "29.80", "Y", ""),
        ("2026-09-10", "식비", "25.40", "Y", "저녁"),
        ("2026-09-10", "숙박", "175.00", "Y", ""),
        ("2026-09-10", "교통", "31.50", "N", "택시, 영수증 분실"),
        ("2026-09-11", "식비", "16.80", "Y", ""),
        ("2026-09-11", "교통", "72.00", "Y", "공항 택시"),
    ]),
    dict(name="최은영", dept="구매팀", city="호치민", country="베트남", cur="VND", rows=[
        ("2026-09-14", "교통", "450000", "Y", "공항 택시"),
        ("2026-09-14", "숙박", "2400000", "Y", ""),
        ("2026-09-14", "식비", "520000", "Y", "도착 후 저녁"),
        ("2026-09-15", "식비", "120000", "Y", "조식"),
        ("2026-09-15", "식비", "260000", "N", "시장 식당 점심"),
        ("2026-09-15", "식비", "880000", "Y", "거래처 저녁"),
        ("2026-09-15", "숙박", "2400000", "Y", ""),
        ("2026-09-15", "교통", "180000", "Y", "그랩"),
        ("2026-09-16", "식비", "135000", "Y", ""),
        ("2026-09-16", "식비", "310000", "Y", ""),
        ("2026-09-16", "숙박", "3200000", "Y", "객실 변경", "c"),
        ("2026-09-16", "기타", "1100000", "Y", "통역 비용", "c"),
        ("2026-09-17", "식비", "110000", "Y", "조식"),
        ("2026-09-17", "식비", "420000", "Y", "공장 방문 점심"),
        ("2026-09-17", "숙박", "2400000", "Y", ""),
        ("2026-09-17", "교통", "230000", "N", "그랩, 앱 내역만 있음"),
        ("2026-09-18", "식비", "150000", "Y", ""),
        ("2026-09-18", "교통", "470000", "Y", "공항 택시"),
        ("2026-09-18", "기타", "480000", "Y", "샘플 포장비"),
    ]),
    dict(name="정혜원", dept="기획팀", city="도쿄", country="일본", cur="JPY", rows=[
        ("2026-09-23", "교통", "3070", "Y", "나리타 익스프레스"),
        ("2026-09-23", "숙박", "18000", "Y", "", "c"),
        ("2026-09-23", "식비", "2400", "Y", "저녁"),
        ("2026-09-24", "식비", "900", "Y", "조식"),
        ("2026-09-24", "식비", "1400", "Y", ""),
        ("2026-09-24", "식비", "4600", "Y", "저녁"),
        ("2026-09-24", "숙박", "18000", "Y", ""),
        ("2026-09-24", "교통", "3000", "Y", "교통카드 충전"),
        ("2026-09-24", "기타", "4800", "N", "선물"),
        ("2026-09-25", "식비", "1000", "Y", ""),
        ("2026-09-25", "식비", "2200", "Y", ""),
        ("2026-09-25", "숙박", "31500", "Y", "연휴 요금", "c"),
        ("2026-09-25", "교통", "1800", "Y", ""),
        ("2026-09-26", "식비", "1800", "Y", ""),
        ("2026-09-26", "식비", "14800", "Y", "지사 직원 저녁"),
        ("2026-09-26", "숙박", "18000", "Y", ""),
        ("2026-09-26", "교통", "2100", "Y", ""),
        ("2026-09-27", "식비", "1500", "N", "공항 점심"),
        ("2026-09-27", "교통", "4200", "Y", "공항 이동"),
    ]),
    dict(name="강태우", dept="연구개발팀", city="프랑크푸르트", country="독일", cur="EUR", rows=[
        ("2026-10-02", "교통", "55.00", "Y", "공항 택시"),
        ("2026-10-02", "숙박", "142.00", "Y", ""),
        ("2026-10-02", "식비", "26.50", "Y", "저녁"),
        ("2026-10-03", "식비", "11.80", "Y", "조식"),
        ("2026-10-03", "식비", "17.60", "Y", ""),
        ("2026-10-03", "식비", "29.50", "Y", "저녁"),
        ("2026-10-03", "숙박", "152.00", "Y", "박람회 기간 요금"),
        ("2026-10-03", "교통", "10.20", "Y", "일일권"),
        ("2026-10-04", "식비", "12.40", "Y", "조식"),
        ("2026-10-04", "식비", "21.00", "Y", ""),
        ("2026-10-04", "식비", "28.80", "Y", "저녁"),
        ("2026-10-04", "숙박", "142.00", "Y", ""),
        ("2026-10-04", "교통", "24.50", "N", "우버"),
        ("2026-10-05", "식비", "13.20", "Y", ""),
        ("2026-10-05", "식비", "19.40", "Y", ""),
        ("2026-10-05", "숙박", "142.00", "Y", ""),
        ("2026-10-05", "기타", "28.00", "Y", "입장권"),
        ("2026-10-06", "기타", "38.00", "Y", "샘플 배송비"),
        ("2026-10-06", "교통", "58.00", "Y", "공항 택시"),
    ]),
    dict(name="윤서진", dept="마케팅팀", city="상하이", country="중국", cur="CNY", rows=[
        ("2026-10-08", "교통", "80.00", "Y", "공항 마그레브"),
        ("2026-10-08", "숙박", "780.00", "Y", ""),
        ("2026-10-08", "식비", "62.00", "Y", ""),
        ("2026-10-08", "식비", "148.00", "Y", "저녁"),
        ("2026-10-09", "식비", "22.00", "N", "조식"),
        ("2026-10-09", "식비", "58.00", "Y", ""),
        ("2026-10-09", "식비", "620.00", "Y", "거래처 저녁"),
        ("2026-10-09", "숙박", "780.00", "Y", ""),
        ("2026-10-09", "교통", "45.00", "Y", "디디 택시"),
        ("2026-10-10", "식비", "30.00", "Y", "조식"),
        ("2026-10-10", "교통", "190.00", "Y", "공항 택시"),
        ("2026-10-10", "기타", "220.00", "N", "기념품"),
    ]),
]

# groups that must exceed the limit after conversion (name, date, item) -- everything else must not
PLANTED = {
    ("박민석", "2026-09-08", "식비"),
    ("박민석", "2026-09-09", "숙박"),
    ("최은영", "2026-09-16", "기타"),
    ("정혜원", "2026-09-25", "숙박"),
    ("정혜원", "2026-09-26", "식비"),
    ("강태우", "2026-10-06", "기타"),
    ("윤서진", "2026-10-09", "식비"),
}


# ---------------------------------------------------------------- rates
def fetch_raw():
    RAW.mkdir(parents=True, exist_ok=True)
    for cur, code in CODES.items():
        page, start = 1, 1
        while True:
            url = (f"https://ecos.bok.or.kr/api/StatisticSearch/sample/json/kr/{start}/{start + 9}"
                   f"/731Y001/D/{RANGE[0]}/{RANGE[1]}/{code}")
            body = urllib.request.urlopen(url, timeout=30).read()
            (RAW / f"{cur}_p{page:02d}.json").write_bytes(body)
            total = json.loads(body)["StatisticSearch"]["list_total_count"]
            start += 10
            page += 1
            if start > total:
                break
            time.sleep(0.3)


def load_rates():
    """{cur: {YYYY-MM-DD: Decimal}} from sample_raw."""
    if "--refetch" in sys.argv or not (RAW / "USD_p01.json").exists():
        fetch_raw()
    rates = {c: {} for c in CODES}
    for cur in CODES:
        for f in sorted(RAW.glob(f"{cur}_p*.json")):
            for r in json.loads(f.read_text(encoding="utf-8"))["StatisticSearch"]["row"]:
                t = r["TIME"]
                rates[cur][f"{t[:4]}-{t[4:6]}-{t[6:]}"] = Decimal(r["DATA_VALUE"])
    return rates


def rate_for(rates, cur, date):
    """(rate per 1 unit of currency in KRW, base date) using the previous business day if needed."""
    days = [d for d in rates[cur] if d <= date]
    base = max(days)
    return rates[cur][base] / UNIT[cur], base


def krw(rates, cur, date, amount):
    r, base = rate_for(rates, cur, date)
    return int((Decimal(amount) * r).quantize(Decimal("1"), rounding=ROUND_HALF_UP)), base


# ---------------------------------------------------------------- checks
def build_rows():
    out = []
    n = 0
    for t in TRIPS:
        for row in t["rows"]:
            n += 1
            d, item, amt, ev, memo = row[:5]
            comma = len(row) > 5
            shown = f"{int(amt):,}" if comma else amt
            out.append([f"X{n:03d}", d, t["name"], t["dept"], t["city"], t["country"], t["cur"], shown, item, ev, memo])
    return out


def verify(rates):
    groups = {}
    for t in TRIPS:
        for d, item, amt, ev, memo, *_ in t["rows"]:
            v, base = krw(rates, t["cur"], d, amt)
            key = (t["name"], d, item)
            groups.setdefault(key, [0, 0, base, t["cur"]])
            g = groups[key]
            g[0] += v
            g[1] += 1
    print("=== over-limit check (per person-day-item group) ===")
    print(f"{'person':6} {'date':10} {'item':3} {'cur':3} {'rows':>4} {'rate base':10} {'KRW total':>11} {'limit':>9}  result")
    over = set()
    near = []
    for key, (tot, cnt, base, cur) in sorted(groups.items(), key=lambda kv: kv[0]):
        item = key[2]
        if item == "교통":
            continue
        lim = LIMIT[item]
        if item == "기타":
            assert cnt == 1, key      # limit is per receipt; one row per group here
        if tot > lim:
            over.add(key)
            print(f"{key[0]:6} {key[1]} {item:3} {cur:3} {cnt:>4} {base:10} {tot:>11,} {lim:>9,}  OVER (+{tot - lim:,})")
        elif tot >= lim * 0.85:
            near.append((key, tot, lim))
    print("--- near-limit groups that stay under (>=85%) ---")
    for key, tot, lim in near:
        print(f"{key[0]:6} {key[1]} {key[2]:3} {tot:>11,} / {lim:>9,}")
    assert over == PLANTED, f"over-limit set mismatch: extra={over - PLANTED}, missing={PLANTED - over}"
    print(f"over-limit groups: {len(over)} (matches planted set)")
    return groups, over


# ---------------------------------------------------------------- texts
RULES = """# 해외 출장 경비 정산 규정

이 규정은 해외 출장에서 쓴 경비를 정산할 때 쓰는 기준이다. 모든 금액은 원화로 환산해서 비교한다.

## 1. 일일 한도

한도는 출장자 한 사람 기준이다. 식비는 하루에 쓴 식비를 모두 더해서, 숙박은 1박마다, 기타는 영수증 1건마다 비교한다.

| 항목 | 한도 | 비교 단위 |
|------|------|-----------|
| 식비 | 100,000원 | 1인 1일 합계 |
| 숙박 | 250,000원 | 1박 |
| 기타 | 50,000원 | 건당 |
| 교통 | 한도 없음 | - |

## 2. 증빙

영수증이나 카드 전표 같은 증빙이 없는 지출은 정산에서 주의 대상으로 표시한다. 증빙이 없다는 이유만으로 지급을 막지는 않지만, 담당자가 사유를 확인한다.

## 3. 판정

- 정상: 한도 안이고 증빙이 있다.
- 주의: 증빙이 없다.
- 초과: 한도를 넘었다. 증빙이 없어도 한도를 넘었으면 초과로 본다.

## 4. 환율

- 환율은 사용일 당일의 한국은행 매매기준율을 쓴다.
- 사용일에 환율이 없으면(주말, 공휴일) 직전 영업일의 환율을 쓴다.
- 엔(JPY)과 동(VND)은 100 단위 환율이므로 환율을 100으로 나눠서 곱한다.
- 원화 금액은 원 단위 미만을 반올림한다.
- 식비처럼 하루 여러 건을 더하는 항목은 건별로 환산해 반올림한 뒤 더한다.
"""

API = """# ECOS 환율 API 안내

한국은행 경제통계시스템(ECOS) 오픈 API로 일별 환율을 가져온다. 주소의 도메인은 ecos.bok.or.kr 이다.

## 요청 주소

```
https://ecos.bok.or.kr/api/StatisticSearch/{인증키}/json/kr/{시작건}/{끝건}/731Y001/D/{시작일}/{종료일}/{항목코드}
```

- 731Y001: 주요국 통화의 대원화환율, D: 일별
- 시작일, 종료일은 YYYYMMDD (예: 20260907)

## 항목코드

| 통화 | 항목코드 | 비고 |
|------|----------|------|
| USD | 0000001 | 1달러 |
| JPY | 0000002 | 100엔 단위 |
| EUR | 0000003 | 1유로 |
| CNY | 0000053 | 1위안 (0000027 은 2016년에 끝난 옛 코드) |
| VND | 0000035 | 100동 단위 |

## 응답 모양

```json
{"StatisticSearch":{"list_total_count":35,"row":[{"STAT_CODE":"731Y001","TIME":"20261001","DATA_VALUE":"1355.7"}]}}
```

- DATA_VALUE 는 문자열이고 천 단위 쉼표가 없다. 숫자로 바꿔서 쓴다.
- 환율이 없는 날(주말, 공휴일)은 row 에 아예 나오지 않는다.
- 기간 안에 데이터가 하나도 없으면 아래처럼 온다.

```json
{"RESULT":{"CODE":"INFO-200","MESSAGE":"해당하는 데이터가 없습니다."}}
```

## 인증키

- 가입 없이 쓸 수 있는 공개 시험 키 `sample` 이 있다. 대신 한 번에 10건까지만 돌려준다.
  1/10, 11/20, 21/30 ... 식으로 시작건/끝건을 바꿔 가며 나눠 받는다.
- 본인 인증키를 받았다면 .env 파일에 ECOS_API_KEY=발급받은키 로 넣고 코드에서 읽어 쓴다.
- 인증키는 코드나 채팅에 그대로 붙여 넣지 않는다.

## 예시 (sample 키, 달러, 10월 첫 열흘)

```
https://ecos.bok.or.kr/api/StatisticSearch/sample/json/kr/1/10/731Y001/D/20261001/20261010/0000001
```
"""

START = f"""# 출장경비 실습키트 시작하기

이 폴더는 해외 출장 경비 정산서를 자동으로 만드는 실습 재료다. 순서대로 열어 보면 된다.

expenses.csv 는 출장 다섯 건의 경비 내역이다. 출장자와 지출 내용은 모두 지어낸 것이라 실제 사람이나 회사와 관련이 없다. 금액은 현지 통화 그대로 들어 있고, 일부는 1,200,000 처럼 쉼표가 섞여 있어서 정리가 필요하다.

정산규정.md 는 한도와 환율 적용 기준이 적힌 회사 규정이고, API_안내.md 는 한국은행 ECOS 에서 환율을 받는 방법이다. 환율은 지어낸 값이 아니라 한국은행이 발표한 실제 매매기준율이다.

정산서_템플릿.xlsx 는 결과를 담을 빈 정산서 양식이다. 머리글과 판정 색깔 규칙만 들어 있다.

.env.example 을 .env 로 복사하면 시험용 키 sample 로 바로 시작할 수 있다. 본인 인증키가 생기면 그 값으로 바꾼다.

sample_raw 폴더에는 {FETCH_DATE} 에 sample 키로 받아 둔 ECOS 실제 응답이 들어 있다. 통화별로 10건씩 나눈 JSON 파일(USD_p01.json 같은 이름)이고, 2026-08-25 부터 2026-10-16 까지의 환율이다. 인증키가 아직 없거나 네트워크가 안 될 때 API 대신 이 파일을 읽어서 쓴다. 직접 받은 값이 아니라 미리 저장해 둔 값이라는 점만 기억하자.
"""

ENV = "# 본인 인증키를 받으면 sample 을 지우고 붙여넣으세요\nECOS_API_KEY=sample\n"


# ---------------------------------------------------------------- xlsx
def make_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.formatting.rule import CellIsRule

    wb = Workbook()
    wb.properties.creator = "ai-agent-class"
    wb.properties.created = dt.datetime(2026, 10, 22)
    wb.properties.modified = dt.datetime(2026, 10, 22)
    ws = wb.active
    ws.title = "정산서"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws["A1"] = "해외 출장 경비 정산서"
    ws["A1"].font = Font(bold=True, size=16)
    for r, label in ((2, "출장자"), (3, "출장지"), (4, "기간")):
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2).border = border
    for r, label in ((2, "총 건수"), (3, "원화 합계"), (4, "정상"), (5, "주의"), (6, "초과")):
        c = ws.cell(r, 8, label)
        c.font = Font(bold=True)
        c.border = border
        ws.cell(r, 9).border = border

    headers = ["전표번호", "사용일", "출장자", "도시", "통화", "금액", "적용환율", "환율기준일",
               "원화금액", "항목", "증빙", "판정", "사유"]
    widths = [10, 12, 10, 14, 7, 14, 11, 12, 14, 8, 7, 8, 30]
    fill = PatternFill("solid", start_color="1F3864", end_color="1F3864")
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(7, i, h)
        c.fill = fill
        c.font = Font(bold=True, color="FFFFFF")
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A8"
    ws["I3"].number_format = "#,##0"
    for r in range(8, 501):
        ws.cell(r, 9).number_format = "#,##0"

    red = PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE")
    yellow = PatternFill("solid", start_color="FFEB9C", end_color="FFEB9C")
    ws.conditional_formatting.add("L8:L500", CellIsRule(operator="equal", formula=['"초과"'], fill=red))
    ws.conditional_formatting.add("L8:L500", CellIsRule(operator="equal", formula=['"주의"'], fill=yellow))

    buf = io.BytesIO()
    wb.save(buf)
    # re-pack with fixed entry timestamps so the bytes are reproducible
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(buf.getvalue())) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for name in sorted(zin.namelist()):
            zi = zipfile.ZipInfo(name, ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            data = zin.read(name)
            if name == "docProps/core.xml":      # openpyxl stamps 'modified' with the current time on save
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-10-22T00:00:00Z", data)
            zout.writestr(zi, data)
    return out.getvalue()


# ---------------------------------------------------------------- main
def main():
    rates = load_rates()
    verify(rates)
    rows = build_rows()

    KIT.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["전표번호", "사용일", "출장자", "부서", "도시", "국가", "통화", "금액", "항목", "증빙", "메모"])
    w.writerows(rows)
    csv_bytes = ("﻿" + buf.getvalue()).encode("utf-8")

    files = {
        "expenses.csv": csv_bytes,
        "정산규정.md": RULES.encode("utf-8"),
        "API_안내.md": API.encode("utf-8"),
        "시작하기.md": START.encode("utf-8"),
        ".env.example": ENV.encode("utf-8"),
        "정산서_템플릿.xlsx": make_xlsx(),
    }
    for name, data in files.items():
        (KIT / name).write_bytes(data)
    for f in sorted(RAW.glob("*.json")):
        files[f"sample_raw/{f.name}"] = f.read_bytes()

    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(files):
            zi = zipfile.ZipInfo(f"{TOP}/{name}", ZIP_TIME)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, files[name])
    print(f"rows: {len(rows)}  zip: {ZIP} ({ZIP.stat().st_size} bytes, {len(files)} files)")


if __name__ == "__main__":
    main()
